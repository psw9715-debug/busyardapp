# -*- coding: utf-8 -*-
"""tools/inbox.py 검사 — 폰이 올린 판을 인쇄용 엑셀에 제대로 채우는가

    python -m unittest tools/test_inbox.py
"""
import datetime
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import inbox  # noqa: E402

GREY = "FF8F8F8F"


def board(entries):
    return {"date": "2026-09-21", "yard": "new", "layout": inbox.current_layout(),
            "updatedAt": "2026-09-21T01:02:03.000Z", "entries": entries}


class WorkDate(unittest.TestCase):
    def test_before_nine_is_previous_day(self):
        self.assertEqual(inbox.work_date(datetime.datetime(2026, 9, 21, 8, 59)), "2026-09-20")

    def test_from_nine_is_same_day(self):
        self.assertEqual(inbox.work_date(datetime.datetime(2026, 9, 21, 9, 0)), "2026-09-21")


class Fill(unittest.TestCase):
    def setUp(self):
        self.cells = inbox.spot_cells()
        self.out = os.path.join(tempfile.mkdtemp(), "out.xlsx")

    def sheet(self):
        wb = openpyxl.load_workbook(self.out)
        self.assertEqual(wb.sheetnames, [inbox.SHEET], "인쇄할 시트만 남긴다")
        return wb[inbox.SHEET]

    def test_spot_cells_cover_route_and_extras(self):
        self.assertEqual(len(self.cells), 196)
        self.assertEqual(self.cells[1], "N3")        # B5-1

    def test_plates_written_vacant_blank(self):
        n = inbox.fill(board({
            "1": {"plate": "1001", "status": "filled", "round": 1},
            "2": {"plate": None, "status": "vacant", "round": 1},
        }), self.out)
        ws = self.sheet()
        self.assertEqual(n, 1)
        self.assertEqual(ws[self.cells[1]].value, "1001")
        self.assertIsNone(ws[self.cells[2]].value, "공차는 종이에서 빈 칸")

    def test_round1_grey_only_when_round2_exists(self):
        inbox.fill(board({"1": {"plate": "1001", "status": "filled", "round": 1}}), self.out)
        font = self.sheet()[self.cells[1]].font
        self.assertNotEqual(font.color.rgb if font.color else None, GREY, "1회차뿐이면 흐리게 하지 않는다")

        inbox.fill(board({
            "1": {"plate": "1001", "status": "filled", "round": 1},
            "3": {"plate": "1503", "status": "filled", "round": 2},
        }), self.out)
        ws = self.sheet()
        self.assertEqual(ws[self.cells[1]].font.color.rgb, GREY)
        self.assertFalse(ws[self.cells[1]].font.b)
        self.assertTrue(ws[self.cells[3]].font.b)
        self.assertNotEqual(ws[self.cells[3]].font.color.rgb if ws[self.cells[3]].font.color else None, GREY)

    def test_layout_mismatch_refused(self):
        b = board({"1": {"plate": "1001", "status": "filled", "round": 1}})
        b["layout"] = inbox.current_layout() - 1
        with self.assertRaises(ValueError):
            inbox.fill(b, self.out)


if __name__ == "__main__":
    unittest.main()
