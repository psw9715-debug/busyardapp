# -*- coding: utf-8 -*-
"""걷는 순서표를 만든다 → src/walk-order.js

자리 번호와 걷는 순서는 같지 않다.

- 신차고지: 쉼터 세 칸(휴게실-1·2·3)은 눈앞에서 나란히 붙어 있는데 번호는
  21 · 132 · 133 으로 멀리 떨어져 있다. 그래서 휴1 을 부른 뒤 휴2·휴3 을 건너뛰고
  B2-1 로 넘어갔다. 셋을 이어 붙인다.
- 구차고지: `차고지 프린트.xlsx` 의 `구차고지-순서` 시트에 실제로 걷는 차례가
  1~89 로 적혀 있다. 그 차례를 그대로 쓴다.

여기서 정하는 것은 **걷는 차례뿐**이다. 자리 번호는 그대로 두므로 저장해 둔 판도,
엑셀 칸 짝짓기도 그대로다.

    py tools\\build_walk_order.py
"""
import io
import json
import os
import re

import openpyxl

뿌리 = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(뿌리, "reference", "차고지 프린트.xlsx")
낼곳 = os.path.join(뿌리, "src", "walk-order.js")


def 자리표(이름):
    """생성된 자리표에서 {엑셀칸: 자리번호} 와 {자리번호: 이름} 을 꺼낸다."""
    s = io.open(os.path.join(뿌리, "src", 이름), encoding="utf-8").read()
    j = json.loads(s[s.index("{"): s.rindex("}") + 1])
    칸 = {c["xl"]: c["spot"] for c in j["cells"] if c.get("kind") == "spot"}
    이름표 = {c["spot"]: c.get("label", "") for c in j["cells"] if c.get("kind") == "spot"}
    딸림 = {c["spot"] for c in j["cells"] if c.get("kind") == "spot" and c.get("extra")}
    return 칸, 이름표, 딸림


# ── 신차고지: 쉼터 셋을 이어 붙인다 ─────────────────────────────────
_, 신이름, 신딸림 = 자리표("yard-data.js")
쉼터 = sorted((n for n, l in 신이름.items() if l.startswith("휴게실")),
              key=lambda n: int(신이름[n].split("-")[1]))
assert len(쉼터) == 3, "신차고지 쉼터가 셋이 아니다: %s" % 쉼터
첫쉼터 = 쉼터[0]
신순서 = {}
for n, 이름 in 신이름.items():
    if n in 신딸림:
        continue
    신순서[n] = float(n)
for 자리, 더함 in zip(쉼터[1:], (0.1, 0.2)):
    신순서[자리] = 첫쉼터 + 더함          # 휴1 바로 뒤로 끌어온다

# ── 구차고지: 시트에 적힌 차례 그대로 ──────────────────────────────
구칸, 구이름, 구딸림 = 자리표("yard-old-data.js")
ws = openpyxl.load_workbook(XLSX, data_only=True)["구차고지-순서"]
구순서 = {}
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is None or not str(v).strip().isdigit():
            continue
        칸이름 = "%s%d" % (openpyxl.utils.get_column_letter(c), r)
        자리 = 구칸.get(칸이름)
        if 자리 is None:
            continue                      # 자리가 아닌 칸에 적힌 숫자는 건너뛴다
        구순서[자리] = float(v)

걸어야할것 = {n for n in 구이름 if n not in 구딸림}
빠진것 = sorted(걸어야할것 - set(구순서))
assert not 빠진것, "구차고지에서 차례를 못 찾은 자리: %s" % [(n, 구이름[n]) for n in 빠진것]
assert len(set(구순서.values())) == len(구순서), "구차고지 차례에 같은 번호가 둘 있다"

본문 = """// 자동 생성 파일 — 직접 수정하지 말 것. `py tools/build_walk_order.py` 로 재생성.
//
// 걷는 순서표. 자리 번호와 걷는 차례는 같지 않다.
//   신차고지 — 쉼터 세 칸은 눈앞에 나란히 있는데 번호가 21·132·133 으로 떨어져 있다.
//   구차고지 — '차고지 프린트.xlsx' 의 `구차고지-순서` 시트에 적힌 차례 그대로다.
// 자리 번호는 건드리지 않는다. 저장해 둔 판도, 엑셀 칸 짝짓기도 그대로다.
export const WALK_ORDER = %s;
""" % json.dumps({"new": {str(k): v for k, v in sorted(신순서.items(), key=lambda x: x[1])},
                  "old": {str(k): v for k, v in sorted(구순서.items(), key=lambda x: x[1])}},
                 ensure_ascii=False, indent=2)

io.open(낼곳, "w", encoding="utf-8", newline="\n").write(본문)

차례 = sorted(신순서, key=lambda n: 신순서[n])
print("신차고지 %d자리 — 쉼터 둘레: %s" % (
    len(신순서), " → ".join(신이름[n] for n in 차례[차례.index(첫쉼터) - 1: 차례.index(첫쉼터) + 4])))
차례2 = sorted(구순서, key=lambda n: 구순서[n])
print("구차고지 %d자리 — 처음 다섯: %s" % (
    len(구순서), " → ".join(구이름[n] for n in 차례2[:5])))
print("낸 곳:", 낼곳)
