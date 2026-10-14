# -*- coding: utf-8 -*-
"""빨라졌는가, 그리고 한 번만 적는가. 원본은 읽기만 하고 복사본에만 쓴다."""
import datetime
import io
import json
import os
import shutil
import sys
import time

sys.path.insert(0, r"D:\프로그램2\꾀꼬리\busyardapp\agent")
sys.path.append(r"D:\프로그램2\꾀꼬리\busyardapp\tools")
import agent  # noqa: E402
import ops_fill  # noqa: E402
import openpyxl  # noqa: E402

agent._data_path = lambda name: os.path.join(r"D:\프로그램2\꾀꼬리\busyardapp\src", name)
agent.log = lambda m: None
agent.notice = lambda *a, **k: None

ok = fail = 0


def eq(got, want, label):
    global ok, fail
    if got == want:
        ok += 1
        print(f"  ok   {label}")
    else:
        fail += 1
        print(f"  FAIL {label}\n       기대 {want!r}\n       실제 {got!r}")


orig = ops_fill.find_excel(datetime.date(2026, 10, 9))
zdir = os.path.join(r"Z:\교통사업처_버스운영센터\상황실", "받아쓰기")
copy = os.path.join(zdir, "속도시험.xlsx")


def fresh():
    if os.path.exists(copy):
        os.remove(copy)
    shutil.copy2(orig, copy)


board = {"yard": "new", "date": "2026-10-08", "entries": {
    str(n): ({"plate": None, "status": "vacant", "round": 1} if n % 9 == 0 else
             {"plate": str(1500 + n), "status": "filled", "round": 1}) for n in range(1, 137)}}

print("▶ A 가 열어 둔 엑셀에 적을 때 (가장 흔한 길)")
fresh()
import pythoncom
import win32com.client
pythoncom.CoInitialize()
holder = win32com.client.DispatchEx("Excel.Application")
holder.Visible = True
holder.DisplayAlerts = False
held = holder.Workbooks.Open(copy)
time.sleep(1)
t = time.time()
written, was_open = agent.write_board(board, copy, log=lambda m: None)
a_open = time.time() - t
print(f"   {written}대 / {a_open:.1f}초")
eq(was_open, True, "열어 둔 그 엑셀에 적었다")
eq(a_open < 12, True, f"열어 둔 엑셀이면 10초 안팎에 끝난다 (잰 값 {a_open:.1f}초)")
eq(bool(holder.EnableEvents), True, "이벤트를 도로 켜 둔다 (남의 엑셀이다)")
eq(holder.Calculation, -4105, "계산 방식도 그대로 되돌린다")
held.Close(SaveChanges=False)
agent.quit_excel(holder)

print("▶ 아무도 안 열어 두었을 때 (A 가 공유 폴더로)")
fresh()
t = time.time()
written, was_open = agent.write_board(board, copy, log=lambda m: None)
a_closed = time.time() - t
print(f"   {written}대 / {a_closed:.1f}초")
eq(was_open, False, "우리가 열어서 적었다")
eq(a_closed < 30, True, f"30초 안에 끝난다 (잰 값 {a_closed:.1f}초)")

print("▶ 사무실 PC 가 적을 때 — 한 번만 적는다")
fresh()
ops_fill.find_excel = lambda date: copy
import inbox
inbox.fetch = lambda date, yard="new": board
t = time.time()
summary = ops_fill.run(board_date="2026-10-08", log=lambda m: None, yard="new")
b_time = time.time() - t
print(f"   {summary} / {b_time:.1f}초")
eq("원본에" in summary, True, "원본에 바로 넣는다 (복사본을 거치지 않는다)")
eq(b_time < 30, True, f"30초 안에 끝난다 (잰 값 {b_time:.1f}초, 전에는 두 벌 넣어 두 배였다)")
eq(os.path.exists(ops_fill.copy_path(copy, datetime.date(2026, 10, 9))), False,
   "쓸데없는 복사본을 만들지 않는다")

after = openpyxl.load_workbook(copy, data_only=True)["고정차고지_입력"]
cells = ops_fill.block_cells("new")
eq(after.cell(*cells["B5-1"]).value, 1501, "B5-1 = 1501")
vacant = ops_fill._label_of(9, "new")      # 9번 자리를 공차로 넣었다
eq(after.cell(*cells[vacant]).value, None, f"공차 자리({vacant})는 비운다")

print("▶ 누가 열어 두었으면 복사본으로 돌아간다")
fresh()
holder = win32com.client.DispatchEx("Excel.Application")
holder.DisplayAlerts = False
held = holder.Workbooks.Open(copy)
time.sleep(1)
summary = ops_fill.run(board_date="2026-10-08", log=lambda m: None, yard="new")
eq("열어 두어" in summary, True, "원본은 건드리지 않고 그 사정을 알린다")
eq("차고지입력" in summary, True, "대신 복사본에 넣는다")
held.Close(SaveChanges=False)
agent.quit_excel(holder)
made = ops_fill.copy_path(copy, datetime.date(2026, 10, 9))
if os.path.exists(made):
    os.remove(made)
os.remove(copy)

print(f"\n   ── 걸린 시간 ──")
print(f"   A 가 열어 둔 엑셀에      {a_open:5.1f}초   (전: 같은 길에서 13초쯤)")
print(f"   A 가 열어서              {a_closed:5.1f}초")
print(f"   사무실 PC 가 원본에      {b_time:5.1f}초   (전: 복사본+원본 두 벌로 70초쯤)")
print(f"\n통과 {ok} / 실패 {fail}")
sys.exit(1 if fail else 0)
