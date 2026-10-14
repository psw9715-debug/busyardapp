# -*- coding: utf-8 -*-
"""1차고지(구차고지)도 A 가 열어 둔 엑셀에 그대로 적는가. 그리고 사람이 쓰고 있을 때
오타·반쯤 적기가 생기지 않는가. 원본은 건드리지 않고 Z: 의 복사본으로만 시험한다."""
import datetime
import os
import shutil
import sys
import time

sys.path.insert(0, r"D:\프로그램2\꾀꼬리\busyardapp\agent")
sys.path.append(r"D:\프로그램2\꾀꼬리\busyardapp\tools")
import agent  # noqa: E402
import openpyxl  # noqa: E402

agent._data_path = lambda name: os.path.join(r"D:\프로그램2\꾀꼬리\busyardapp\src", name)
agent.log = lambda m: None
agent.notice = lambda *a, **k: None          # 시험 중에는 창을 띄우지 않는다

ok = fail = 0


def eq(got, want, label):
    global ok, fail
    if got == want:
        ok += 1
        print(f"  ok   {label}")
    else:
        fail += 1
        print(f"  FAIL {label}\n       기대 {want!r}\n       실제 {got!r}")


src = None
for d in range(0, 40):
    try:
        src = agent.find_excel(datetime.date(2026, 9, 26) + datetime.timedelta(days=d))
        break
    except FileNotFoundError:
        continue
zdir = os.path.join(r"Z:\교통사업처_버스운영센터\상황실", "받아쓰기")
os.makedirs(zdir, exist_ok=True)
copy = os.path.join(zdir, "1차고지시험.xlsx")
if os.path.exists(copy):
    os.remove(copy)
shutil.copy2(src, copy)
print("   본보기:", os.path.basename(src))

old = agent.block_cells("old")
new = agent.block_cells("new")
print("▶ 1차고지 자리표")
eq(len(old), 90, "90자리")
eq(max(r for r, _ in old.values()) <= 34, True, "1~34행 안에만 있다")
eq(min(r for r, _ in new.values()) >= 38, True, "6차고지는 38행부터 — 서로 겹치지 않는다")

# 1차고지 판: 번호 · 공차 · 승용차(전화번호) 를 섞는다
board = {"yard": "old", "date": "2026-10-08", "entries": {}}
for n in range(1, 91):
    if n % 10 == 0:
        board["entries"][str(n)] = {"plate": None, "status": "vacant", "round": 1}
    elif n % 17 == 0:
        board["entries"][str(n)] = {"plate": None, "phone": "01012345678", "status": "car", "round": 1}
    else:
        board["entries"][str(n)] = {"plate": str(3000 + n), "status": "filled", "round": 1}

# 적기 전 6차고지 블록을 적어 둔다 — 그대로여야 한다 (본보기 파일에 이미 값이 있다)
before_new = {k: openpyxl.load_workbook(copy, data_only=True)["고정차고지_입력"].cell(*rc).value
              for k, rc in list(new.items())[:20]}

import pythoncom
import win32com.client
pythoncom.CoInitialize()
holder = win32com.client.DispatchEx("Excel.Application")     # A 가 보고 있는 엑셀
holder.Visible = True
holder.DisplayAlerts = False
held = holder.Workbooks.Open(copy)
time.sleep(1)

print("▶ 1차고지를 열어 둔 엑셀에 적는다")
t = time.time()
written, was_open = agent.write_board(board, copy, log=lambda m: None)
secs = time.time() - t
eq(was_open, True, "열어 둔 그 엑셀에 적었다")
print(f"   {written}대 / {secs:.0f}초")

after = openpyxl.load_workbook(copy, data_only=True)["고정차고지_입력"]
eq(after.cell(*old[1]).value, 3001, "1차고지 첫 자리에 3001")
eq(after.cell(*old[10]).value, None, "공차 자리는 비운다")
eq(after.cell(*old[17]).value, None, "승용차 자리는 비운다 (전화번호를 적지 않는다)")
got = sum(1 for rc in old.values() if after.cell(*rc).value is not None)
eq(got, written, "적었다고 한 대수가 그대로 저장돼 있다")
eq(after.cell(*old[1][0:1] + (old[1][1],)).value, 3001, "저장된 값 그대로")
name_row = after.cell(old[1][0] + 1, old[1][1]).value
eq(name_row not in (None, ""), True, "아래 이름 수식이 살아 값이 찍혔다")

print("▶ 6차고지 칸은 건드리지 않았다")
eq({k: after.cell(*new[k]).value for k in before_new}, before_new,
   "6차고지 블록은 한 칸도 바뀌지 않았다")

print("▶ 사람이 셀을 편집하는 중이라면")
class Busy:
    """엑셀이 '지금은 못 받는다' 고 되돌리는 상황을 흉내 낸다"""
    def __init__(self, n):
        self.left = n
        self.calls = 0

    def __call__(self):
        self.calls += 1
        if self.left > 0:
            self.left -= 1
            raise pythoncom.com_error(-2147417846, "Call was rejected by callee.", None, None)
        return "됐다"

b = Busy(2)
eq(agent.wait_ready(b, "시험", 10, log=lambda m: None), "됐다", "바쁘면 기다렸다 다시 해서 끝낸다")
eq(b.calls, 3, "세 번째에 들어갔다")

b2 = Busy(99)
try:
    agent.wait_ready(b2, "시험", 3, log=lambda m: None)
    eq("안 터졌다", "터져야 한다", "끝내 안 받으면 그대로 알린다")
except pythoncom.com_error:
    eq(True, True, "끝내 안 받으면 그대로 알린다 (몰래 넘기지 않는다)")

print("▶ 계산 방식은 쓰시던 대로 되돌린다")
holder.Calculation = -4135                       # 사람이 '수동' 으로 쓰고 있었다
agent.write_board(board, copy, log=lambda m: None)
eq(holder.Calculation, -4135, "수동으로 쓰고 계셨으면 수동 그대로 둔다")
holder.Calculation = -4105

held.Close(SaveChanges=False)
agent.quit_excel(holder)

print("▶ 아무도 열어 두지 않았다면 (네트워크로 직접)")
import importlib
agent.notice = lambda *a, **k: None
written2, was_open2 = agent.write_board(board, copy, log=lambda m: None)
eq(was_open2, False, "열려 있지 않으면 우리가 열어서 적는다")
eq(written2, written, "같은 대수를 적는다")
after2 = openpyxl.load_workbook(copy, data_only=True)["고정차고지_입력"]
eq(after2.cell(*old[1]).value, 3001, "1차고지 첫 자리에 3001 (닫혀 있어도 그대로)")
eq(len(agent.unc(copy)) > 0 and agent.unc(copy).startswith("\\\\"), True,
   "공유 폴더 경로로 다룬다")

print("▶ 안내 창은 일을 붙잡지 않는다")
importlib.reload(agent)       # notice 를 원래대로 되돌린다
agent._data_path = lambda name: os.path.join(r"D:\프로그램2\꾀꼬리usyardapp\src", name)
agent.log = lambda m: None
t0 = time.time()
th = agent.notice("시험용 안내입니다 (2초 뒤 저절로 닫힙니다)", 2)
eq(time.time() - t0 < 0.5, True, "창을 띄우고 곧바로 일을 이어 간다")
th.join(5)
eq(th.is_alive(), False, "창은 저절로 닫힌다")

os.remove(copy)
print(f"\n통과 {ok} / 실패 {fail}")
sys.exit(1 if fail else 0)
