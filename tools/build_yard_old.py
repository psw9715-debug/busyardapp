# -*- coding: utf-8 -*-
"""구차고지 엑셀 -> src/yard-old-data.js

`차고지 프린트.xlsx` 의 두 시트를 합쳐 하나로 만든다.
- `구차고지-순서` : 자리 이름(1-3-1, 조-1, 한노-1)이 적힌 시트. 자리와 칸 모양은 여기가 정한다.
- `구차고지`      : 실제로 손에 쥐는 인쇄용 시트. 색과 테두리를 여기서 가져온다.

자리 한 칸 = 세 줄(차량번호 / 이름 / 출근시각)을 병합한 것으로, 신차고지와 같다.
다만 블록 간격이 일정하지 않아(제목 2줄, 구분 4줄) 산술로 나눌 수 없다.
그래서 **엑셀에 실제로 있는 블록 시작행을 모아 차례로 격자 줄 번호를 매긴다.**

  그리드 열 = 엑셀 열 번호 (1..15, A..O)
"""
import json
import re

import openpyxl
from openpyxl.utils import get_column_letter as col_letter

from xlsx_color import load_theme, cell_bg, border_weights

XLSX = "reference/차고지 프린트.xlsx"
OUT = "src/yard-old-data.js"
ORDER_SHEET = "구차고지-순서"
PRINT_SHEET = "구차고지"

# 사무실을 나와 도는 순서. 한 덩어리가 "다음" 으로 건너뛰는 단위다.
ROUTE = [
    ("1-3", 1, 13), ("2-1", 1, 9), ("2-2", 1, 10), ("2-3", 1, 9),
    ("조", 1, 11), ("한노", 1, 8), ("1-1", 1, 15), ("1-2", 1, 14),
]
SPOT_LABEL = re.compile(r"^(\d-\d|조|한노)-(\d+)$")
EXTRA_LABEL = "예비"


def grid_rows(ws):
    """엑셀 행 -> 격자 줄 번호. 병합 블록이 시작하는 행들을 차례로 센다."""
    starts = sorted({r.min_row for r in ws.merged_cells.ranges})
    return {row: i + 1 for i, row in enumerate(starts)}, starts


def span(rows, starts, first, last):
    """엑셀 행 범위가 격자에서 몇 줄을 차지하는가"""
    return max(1, sum(1 for s in starts if first <= s <= last))


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws, wp = wb[ORDER_SHEET], wb[PRINT_SHEET]
    theme = load_theme(XLSX)
    rows, starts = grid_rows(ws)

    order = {}   # "1-3-10" -> (순회 순번, 구간 번호)
    for seg, (zone, first, last) in enumerate(ROUTE):
        for num in range(first, last + 1):
            order[f"{zone}-{num}"] = (len(order) + 1, seg)

    cells = []
    extra = 0
    for rng in sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col)):
        if rng.min_row not in rows:
            continue
        value = ws.cell(rng.min_row, rng.min_col).value
        cell = {
            "col": rng.min_col,
            "colspan": rng.max_col - rng.min_col + 1,
            "row": rows[rng.min_row],
            "rowspan": span(rows, starts, rng.min_row, rng.max_row),
            "xl": col_letter(rng.min_col) + str(rng.min_row),
        }
        bg = cell_bg(wp.cell(rng.min_row, rng.min_col), theme)
        if bg:
            cell["bg"] = bg
        cell["b"] = border_weights(wp, rng.min_row, rng.min_col, rng.max_row, rng.max_col)

        label = str(value).strip() if value is not None else ""
        if SPOT_LABEL.match(label):
            assert label in order, f"{cell['xl']} '{label}' 이(가) ROUTE 에 없다"
            cell["kind"] = "spot"
            cell["spot"], cell["seg"] = order[label]
            cell["label"] = label
        elif label == EXTRA_LABEL:
            # 순서에는 없지만 차를 댈 수 있는 칸. 눌러서 손으로 적는다.
            extra += 1
            cell["kind"] = "spot"
            cell["spot"] = len(order) + extra
            cell["label"] = f"예비-{extra}"
            cell["extra"] = True
        elif value is not None:
            cell["kind"] = "label"
            cell["text"] = str(value).replace("\n", "")
        else:
            cell["kind"] = "void"
        cells.append(cell)

    # 병합에 속하지 않은 칸도 종이에는 선과 색이 있다. 하나라도 빠지면 인쇄물이 달라진다.
    taken = set()
    for c in cells:
        for dr in range(c["rowspan"]):
            for dc in range(c["colspan"]):
                taken.add((c["row"] + dr, c["col"] + dc))

    for excel_row in starts:
        row = rows[excel_row]
        for col in range(1, 16):
            if (row, col) in taken:
                continue
            v = wp.cell(excel_row, col).value
            bg = cell_bg(wp.cell(excel_row, col), theme)
            borders = border_weights(wp, excel_row, col, excel_row + 2, col)
            if v is None and not bg and not any(borders):
                continue
            c = {
                "col": col, "colspan": 1, "row": row, "rowspan": 1,
                "xl": col_letter(col) + str(excel_row),
                "kind": "paint" if v is not None else "plain",
                "b": borders,
            }
            if v is not None:
                c["text"] = str(v)
            if bg:
                c["bg"] = bg
            cells.append(c)

    cells.sort(key=lambda c: (c["row"], c["col"]))
    spots = sorted(c["spot"] for c in cells if c["kind"] == "spot")
    assert spots == list(range(1, len(order) + extra + 1)), "엑셀에 없거나 두 번 적힌 자리가 있다"

    data = {
        "id": "old",
        "name": "구차고지",   # 인쇄 시트 A1 은 자리 라벨이라 이름으로 쓸 수 없다
        "cols": 15,
        "rows": len(starts),
        "wideCol": 1,
        "totalSpots": len(spots),
        "routeSpots": len(order),
        "cells": cells,
    }
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("// 자동 생성 파일 — 직접 수정하지 말 것. `python tools/build_yard_old.py` 로 재생성.\n")
        f.write("export const YARD_OLD = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n")

    kinds = {}
    for c in cells:
        kinds[c["kind"]] = kinds.get(c["kind"], 0) + 1
    print(f"{OUT} 생성 완료 — 자리 {len(spots)} (걷는 순서 {len(order)} + 예비 {extra}), 칸 {kinds}")


if __name__ == "__main__":
    main()
