# -*- coding: utf-8 -*-
"""Z: 의 입출차 운영관리 엑셀 -> guide/data/<날짜>.json -> 깃허브 push

**출차시각은 올리지 않는다.** 저장소가 공개라 시각표가 그대로 인터넷에
남는 것을 피하려는 것이다. 앱이 자리를 정하는 데 필요한 것은 셋뿐이라
그 셋만 한 글자로 줄여 올린다.

    r  휴차 (또는 낮에 나가는 차)   -> 1차고지 맨 뒷열 / 2차고지 2열
    l  6시 15분 이후에 나가는 늦은 차 -> 2차고지 2열
    e  6시까지 나가는 빠른 차        -> 2차고지 1열

1차고지는 애초에 시각을 보지 않으므로 `r` 인지 아닌지만 알면 된다.

폰은 사내 Z: 에 닿을 수 없다. 그래서 Z: 를 읽을 수 있는 사무실 PC 가 다리 역할을 한다.

읽는 날짜는 **근무일 + 1일**이다. 밤에 들어오는 차는 다음날 아침 나가는 차라서다.
근무일은 순회 앱과 같이 오전 9시를 기준으로 넘긴다 (새벽 00:40 도 전날 근무).

폴더·파일 이름은 해마다 달마다 바뀌지만 시트 이름 `차고지참고` 는 그대로다.
그래서 이름은 전부 glob 으로 찾고 시트만 이름으로 집는다.

    python tools/push_source.py            # 근무일+1
    python tools/push_source.py 2026-09-10 # 날짜를 직접
    python tools/push_source.py --no-push  # 파일만 만들고 push 안 함
"""
import collections
import datetime
import glob
import json
import os
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = r"Z:\교통사업처_버스운영센터\상황실"
SHEET = "차고지참고"
OUT_DIR = os.path.join(ROOT, "guide", "data")

DAY_OUT = "09:00"     # 이 시각 이후에 나가면 사실상 안 나가는 것으로 본다
LANE2_ROOM = 10       # 2차고지 2열 9칸 + 예비 2-27

# 3열이 한 묶음(차량번호 / 출차시각·휴차 / 구역)이고 묶음이 가로로 7개 늘어서 있다
FIRST_COLS = (2, 5, 8, 11, 14, 17, 20)
ROWS = range(3, 48)


def work_date(now=None):
    """오전 9시를 기준으로 날짜를 넘긴다 — 순회 앱 store.js 와 같은 규칙"""
    now = now or datetime.datetime.now()
    d = now.date()
    return d - datetime.timedelta(days=1) if now.hour < 9 else d


def find_one(pattern, what):
    hits = [p for p in glob.glob(pattern) if not os.path.basename(p).startswith("~$")]
    if not hits:
        raise FileNotFoundError(f"{what} 없음 — {pattern}")
    return sorted(hits)[-1]


def locate(date):
    """<연>/<월>/<MMDD>/<MMDD ...입출차 운영관리...>.xlsx 를 찾는다"""
    year = find_one(os.path.join(BASE, f"* {date:%Y}"), "연 폴더")
    month = find_one(os.path.join(year, f"* {date:%Y%m}"), "월 폴더")
    day = find_one(os.path.join(month, f"{date:%m%d}*"), f"{date:%m%d} 폴더")
    return find_one(os.path.join(day, f"{date:%m%d} *입출차 운영관리*.xlsx"), "엑셀")


def parse(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise KeyError(f"'{SHEET}' 시트가 없다 — {wb.sheetnames}")
    grid = list(wb[SHEET].iter_rows(min_row=1, max_row=47, max_col=22, values_only=True))

    raw = {}
    for c0 in FIRST_COLS:
        for r in ROWS:
            plate = grid[r - 1][c0 - 1]
            if plate is None:
                continue
            plate = str(plate).strip()
            if not (plate.isdigit() and len(plate) == 4):
                continue                      # 중간에 끼어 있는 구분 글자
            when = grid[r - 1][c0]
            out = when.strftime("%H:%M") if isinstance(when, (datetime.time, datetime.datetime)) else None
            raw[plate] = None if out is None or out >= DAY_OUT else out
    if not raw:
        raise ValueError("차량을 하나도 못 읽었다 — 시트 모양이 바뀌었는지 확인")
    return raw


def is_yard2(plate):
    n = int(plate)
    if n in (1015, 1016):
        return False                          # 61번은 1차고지다
    return 1012 <= n <= 1031 or 1732 <= n <= 1754


def cutoff_of(raw):
    """2열 정원(10대)에 맞는 컷오프. 실측 4일 모두 06:15 였다."""
    target = [p for p in raw if is_yard2(p)]
    rest = sum(1 for p in target if raw[p] is None)
    room = LANE2_ROOM - rest
    outs = sorted((raw[p] for p in target if raw[p]), reverse=True)
    if room <= 0 or room > len(outs):
        return "06:15"
    return outs[room - 1]


def to_bands(raw):
    """시각을 빼고 r / l / e 세 글자로 줄인다."""
    cut = cutoff_of(raw)
    return cut, {p: ("r" if v is None else ("l" if v >= cut else "e")) for p, v in raw.items()}


def git(*args):
    return subprocess.run(("git",) + args, cwd=ROOT, capture_output=True, text=True)


def main():
    argv = [a for a in sys.argv[1:] if a != "--no-push"]
    push = "--no-push" not in sys.argv

    if argv:
        date = datetime.date.fromisoformat(argv[0])
    else:
        date = work_date() + datetime.timedelta(days=1)

    try:
        path = locate(date)
    except FileNotFoundError as e:
        print(f"[아직 안 올라옴] {date} 자료를 못 찾았다.\n  {e}")
        return 0                              # 오류가 아니다. 나중에 다시 돌리면 된다

    raw = parse(path)
    cut, bands = to_bands(raw)

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, f"{date}.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump({
            "date": str(date),
            "builtAt": datetime.datetime.now().isoformat(timespec="seconds"),
            "cutoff": cut,
            "bands": bands,
        }, f, ensure_ascii=False, separators=(",", ":"))

    n = collections.Counter(bands.values())
    print(f"{out} — {len(bands)}대  휴차 {n['r']} · 늦은 차 {n['l']} · 빠른 차 {n['e']}  (컷오프 {cut})")
    print(f"  원본 {os.path.basename(path)}")
    print("  출차시각은 올리지 않는다 — 세 글자로만 줄여 담았다")

    if not push:
        return 0

    git("add", os.path.relpath(out, ROOT))
    if not git("diff", "--cached", "--quiet").returncode:
        print("  바뀐 것이 없어 push 안 함")
        return 0
    c = git("commit", "-m", f"소스 {date}")
    if c.returncode:
        print("  commit 실패:", c.stderr.strip())
        return 1
    p = git("push")
    if p.returncode:
        print("  push 실패:", p.stderr.strip())
        return 1
    print("  push 완료")
    return 0


if __name__ == "__main__":
    sys.exit(main())
